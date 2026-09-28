"""Default-risk analyzer: model integrity, independent recomputation of the scores, workbook output and API."""
import io
import math

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from fintea.excel import soffice_available, verify_workbook, write_workbook
from fintea.main import app
from fintea.providers import load_dataset
from fintea.risk import RISK_SHEET_ORDER, build_risk_model, derive_inputs, naive_distance_to_default, solve_merton
from fintea.risk.builder import ALT, DASH, DIST, DQ, DUP, MER, PIO, RTG, stamp_verification
from fintea.risk.merton import norm_cdf
from fintea.risk import spec as S

M = 1e6
client = TestClient(app)


@pytest.fixture(scope="module", params=["MSFT", "AAPL", "BYND", "HDFCBANK.NS", "WOLF"])
def risk(request):
    ds = load_dataset(request.param, "sample")
    return build_risk_model(ds)


def test_structure_and_no_errors(risk):
    b = risk.book
    assert b.order == RISK_SHEET_ORDER
    assert risk.summary["error_cells"] == {}
    s = risk.summary
    assert 0 <= s["composite_score"] <= 100 and s["composite_grade"] in {g for _, g in S.GRADES}
    for k in ("pd_merton_naive", "pd_merton_rn", "pd_merton_phys", "zmijewski_pd"):
        assert 0 <= s[k] <= 1, k
    if s["financial_sector"]:
        assert s["pd_rating_1y"] is None and s["synthetic_rating"].startswith("n/a")
    else:
        assert 0 <= s["pd_rating_1y"] <= 1 and 0 <= s["pd_rating_5y"] <= 1 and s["synthetic_rating"] in S.RATING_ORDER
    assert s["merton_check"] < 1e-6
    # every model output is a formula, never a pasted number (except the two solver outputs and the price table)
    for sheet in (ALT, PIO, DIST, RTG, DASH):
        for (r, c), cell in b.sheets[sheet].cells.items():
            if cell.key and isinstance(cell.content, (int, float)) and not isinstance(cell.content, bool):
                assert cell.style == "input", f"{sheet}!{cell.address} ({cell.key}) is a hard-coded non-input"


def test_scores_match_independent_recomputation(risk):
    ds, b = risk.dataset, risk.book
    L = len(ds.periods) - 1
    f = {k: (v or 0.0) / M for k, v in ds.periods[L].fields.items()}
    ta = f["total_assets"]
    mve = ds.market.price * ds.market.shares_outstanding / M
    wc = f["current_assets"] - f["current_liabilities"]
    x1, x2, x3, x4m, x4b, x5 = wc / ta, f["retained_earnings"] / ta, f["operating_income"] / ta, mve / f["total_liabilities"], f["total_equity"] / f["total_liabilities"], f["revenue"] / ta
    z = 1.2 * x1 + 1.4 * x2 + 3.3 * x3 + 0.6 * x4m + 1.0 * x5
    z2 = 6.56 * x1 + 3.26 * x2 + 6.72 * x3 + 1.05 * x4b
    assert abs(b.num(ALT, "z", L) - z) < 1e-9
    assert abs(b.num(ALT, "z2", L) - z2) < 1e-9
    assert b.val(ALT, "z2_zone", L) == ("Safe" if z2 > 2.6 else "Distress" if z2 < 1.1 else "Grey")
    # Zmijewski probit
    x = -4.336 - 4.513 * f["net_income"] / ta + 5.679 * f["total_liabilities"] / ta + 0.004 * (f["current_assets"] / f["current_liabilities"] if f["current_liabilities"] else 0)
    assert abs(b.num(DIST, "x_pd", L) - norm_cdf(x)) < 1e-9
    # Springate
    s_ = 1.03 * x1 + 3.07 * x3 + 0.66 * (f["pretax_income"] / f["current_liabilities"] if f["current_liabilities"] else 0) + 0.4 * x5
    assert abs(b.num(DIST, "s_score", L) - s_) < 1e-9
    # naive distance to default reproduces the Bharath-Shumway closed form
    R = risk.inputs
    E = R.values["price"] * R.values["shares_outstanding"]
    F = f["short_term_debt"] + 0.5 * f["long_term_debt"]
    sigma_e = b.num(MER, "sigma_E")
    mu = b.num(MER, "mu")
    nd = naive_distance_to_default(E, max(F, 1e-6), sigma_e, mu, 1.0)
    assert abs(b.num(MER, "dd_naive") - nd["dd"]) < 1e-7
    assert abs(b.num(MER, "pd_naive") - nd["pd"]) < 1e-9
    # iterated Merton: the workbook formulas reproduce the inputs from the solved (V, sigma_V)
    assert abs(b.num(MER, "chk_E")) < 1e-6 and abs(b.num(MER, "chk_vol")) < 1e-6
    if risk.merton:
        sol = solve_merton(E, sigma_e, F, R.values["risk_free"], 1.0)
        assert abs(sol.asset_value - b.num(MER, "V")) < 1e-6 * E
        assert abs(norm_cdf(-sol.d2) - b.num(MER, "pd_rn")) < 1e-9
    # composite is the weighted average of the sub-scores with the weights actually used
    subs = risk.summary["sub_scores"]
    ws = {k: b.num(DASH, f"{k}_w") for k in subs}
    total_w = sum(ws.values())
    assert abs(b.num(DASH, "composite_score") - sum(subs[k] * ws[k] for k in subs) / total_w) < 1e-9


def test_dupont_identities(risk):
    ds, b = risk.dataset, risk.book
    for p, per in enumerate(ds.periods):
        f = {k: (v or 0.0) / M for k, v in per.fields.items()}
        assert abs(b.num(DUP, "d_roa", p) - f["net_income"] / f["total_assets"]) < 1e-12
        if f["total_equity"] > 0:
            roe = f["net_income"] / f["total_equity"]
            assert abs(b.num(DUP, "d_roe", p) - roe) < 1e-9 and abs(b.num(DUP, "d_roe_check", p) - roe) < 1e-12
            if f["pretax_income"] != 0 and f["operating_income"] != 0:
                assert abs(b.num(DUP, "d_roe5", p) - roe) < 1e-9
        else:
            assert not b.has(DUP, "d_roe", p) and not b.has(DUP, "d_equity_mult", p)
    d = risk.summary["dupont"]
    assert set(d) >= {"roe", "net_margin", "asset_turnover", "equity_multiplier", "tax_burden", "interest_burden", "ebit_margin", "roa"}
    assert any(c["id"] == "dupont_roe" for c in risk.book.to_json()["sheets"][risk.book.order.index(DUP)]["charts"])


def test_single_year_and_financial_cases():
    wolf = build_risk_model(load_dataset("WOLF", "sample"))
    assert wolf.summary["nh"] == 1 and wolf.summary["has_prior_year"] is False
    assert wolf.summary["piotroski"] is None and wolf.summary["beneish_m"] is None and wolf.summary["ohlson_pd"] is None
    assert wolf.book.num(DASH, "sig_pio_w") == 0 and wolf.book.num(DASH, "sig_ohlson_w") == 0
    assert wolf.summary["error_cells"] == {} and wolf.summary["composite_score"] > 50   # negative equity, heavy debt
    bank = build_risk_model(load_dataset("HDFCBANK.NS", "sample"))
    assert bank.summary["financial_sector"] is True
    assert bank.inputs.values["w_merton"] == 1.0 and bank.inputs.values["w_z2"] == 0.0
    assert abs(bank.summary["composite_score"] - bank.summary["sub_scores"]["sig_merton"]) < 1e-9
    assert bank.summary["default_spread"] is None and bank.summary["rating_source"].startswith("Not applicable")
    assert wolf.feedback["distress_votes"] and all(v["model"] != "Ohlson" for v in wolf.feedback["distress_votes"])
    assert any("financial institution" in p.lower() for sec in bank.feedback["qualitative"] for p in sec["points"])


def test_overrides_and_validation():
    ds = load_dataset("MSFT", "sample")
    base = build_risk_model(ds)
    hi = build_risk_model(ds, overrides={"equity_vol_method": 2, "equity_vol_manual": 1.5, "mu_method": 2})
    assert hi.summary["pd_merton_naive"] > base.summary["pd_merton_naive"]
    assert "equity_vol_manual" in hi.inputs.overridden and hi.inputs.basis["equity_vol_manual"].startswith("User override")
    tl = build_risk_model(ds, overrides={"default_point_method": 3})
    assert tl.summary["default_point"] > base.summary["default_point"]
    with pytest.raises(ValueError):
        derive_inputs(ds, {"horizon": 50})
    with pytest.raises(ValueError):
        derive_inputs(ds, {"w_merton": 1.5})


def test_workbook_output(risk):
    data = write_workbook(risk.book)
    assert data[:2] == b"PK"
    wb = load_workbook(io.BytesIO(data))
    assert wb.sheetnames == RISK_SHEET_ORDER
    assert len(wb[DASH]._charts) == len(risk.book.sheets[DASH].charts) >= 8
    for nm in ("CompositeRiskScore", "NaiveDefaultProbability", "SyntheticRating"):
        assert nm in wb.defined_names
    r, c = risk.book.lookup(DASH, "composite_score", None)
    assert str(wb[DASH].cell(row=r, column=c).value).startswith("=")
    stamp_verification(risk, {"status": "verified", "cells_checked": 123, "engine": "LibreOffice Calc (headless)", "max_abs_diff": 1e-9})
    r, c = risk.book.lookup(DQ, "verification_stamp", None)
    assert "123" in risk.book.value(DQ, r, c)


def test_chart_json_resolves_values(risk):
    js = risk.book.to_json()
    dash = next(s for s in js["sheets"] if s["name"] == DASH)
    ids = {c["id"] for c in dash["charts"]}
    assert {"z2_trend", "pd_models", "radar", "leverage", "liquidity"} <= ids
    z2 = next(c for c in dash["charts"] if c["id"] == "z2_trend")
    assert z2["categories"] == risk.summary["labels"]
    assert len(z2["series"][0]["values"]) == risk.summary["nh"]
    assert abs(z2["series"][0]["values"][-1] - risk.summary["altman_z2"]) < 1e-9


@pytest.mark.skipif(not soffice_available(), reason="LibreOffice not installed")
def test_libreoffice_recalculation_matches_engine(risk):
    data = write_workbook(risk.book)
    v = verify_workbook(risk.book, data)
    assert v["status"] == "verified", v["mismatches"][:5]
    assert v["cells_checked"] > 400


def test_api_risk_endpoints():
    r = client.post("/api/risk", json={"query": "BYND", "provider": "sample", "verify": False})
    assert r.status_code == 200, r.text
    js = r.json()
    assert js["kind"] == "risk" and js["summary"]["symbol"] == "BYND" and js["summary"]["composite_grade"] in ("High", "Severe")
    assert js["summary"]["agreement"] >= 3 and js["summary"]["stress"]["pd_both"] >= js["summary"]["pd_merton_naive"] * 0.9
    assert js["summary"]["applicability"] and all(a["status"] for a in js["summary"]["applicability"])
    assert [s["name"] for s in js["sheets"]] == RISK_SHEET_ORDER
    assert js["charts"] and js["inputs"]["items"] and js["feedback"]["qualitative"]
    rid = js["id"]
    # under the risk-free drift a lower equity volatility must lower the naive probability of default
    r2 = client.post(f"/api/risk/{rid}/rebuild", json={"overrides": {"equity_vol_method": 2, "equity_vol_manual": 0.3, "mu_method": 2}, "verify": False, "include_sheets": False})
    assert r2.status_code == 200 and "sheets" not in r2.json()
    r_hi = client.post(f"/api/risk/{rid}/rebuild", json={"overrides": {"equity_vol_method": 2, "equity_vol_manual": 1.5, "mu_method": 2}, "verify": False, "include_sheets": False})
    assert r2.json()["summary"]["pd_merton_naive"] < r_hi.json()["summary"]["pd_merton_naive"]
    assert r2.json()["inputs"]["items"][0]["kind"] == "scalar"
    d = client.get(f"/api/risk/{rid}/download")
    assert d.status_code == 200 and d.content[:2] == b"PK" and "default_risk" in d.headers["content-disposition"]
    assert client.post(f"/api/risk/{rid}/rebuild", json={"overrides": {"horizon": 99}}).status_code == 422
    assert client.get("/api/risk/nope").status_code == 404
