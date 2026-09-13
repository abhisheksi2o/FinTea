import pytest

from fintea.excel import soffice_available, verify_workbook, write_workbook
from fintea.model import build_model, derive_assumptions
from fintea.model.builder import BS, CF, DCF, IS, WACC, BETA, FEED, SHEET_ORDER

TOL = 1e-6


def test_model_integrity(dataset):
    res = build_model(dataset, years=5)
    b = res.book
    assert b.order == SHEET_ORDER
    assert res.summary["error_cells"] == {}
    nh, N = b.meta["nh"], b.meta["np"]
    for p in range(nh + N):
        assert abs(b.num(BS, "balance_check", p)) < TOL, f"balance sheet does not balance in period {p}"
    for p in range(nh, nh + N):
        assert abs(b.num(CF, "cash_check", p)) < TOL
        assert abs(b.num(CF, "end_cash", p) - b.num(BS, "cash", p)) < TOL
    L = nh - 1
    # historical statements reconcile exactly to reported figures
    assert abs(b.num(IS, "ebit", L) - dataset.periods[L].fields["operating_income"] / 1e6) < TOL
    assert abs(b.num(IS, "net_income", L) - dataset.periods[L].fields["net_income"] / 1e6) < TOL
    assert abs(b.num(BS, "total_assets", L) - dataset.periods[L].fields["total_assets"] / 1e6) < TOL
    assert abs(b.num(CF, "cfo", L) - dataset.periods[L].fields["cfo"] / 1e6) < TOL
    # beta cross check and WACC composition
    assert abs(b.num(BETA, "raw_beta") - b.num(BETA, "beta_check")) < 1e-9
    ke = b.num(WACC, "cost_of_equity")
    assert abs(ke - (res.assumptions.values["risk_free"] + b.num(BETA, "selected_beta") * res.assumptions.values["erp"])) < 1e-12
    # DCF identities
    assert abs(b.num(DCF, "enterprise_value") - (b.num(DCF, "sum_pv") + b.num(DCF, "pv_tv"))) < TOL
    assert abs(b.num(DCF, "equity_value") / b.num(DCF, "d_shares") - b.num(DCF, "implied_price")) < TOL
    assert b.num(FEED, "n_fail") == 0
    assert res.feedback["quantitative"] and all(q["status"] in ("PASS", "FLAG", "FAIL") for q in res.feedback["quantitative"])
    assert len(res.feedback["qualitative"]) >= 5


def test_every_projection_is_a_formula(dataset):
    res = build_model(dataset, years=4)
    b = res.book
    nh, N = b.meta["nh"], b.meta["np"]
    for sheet in (IS, BS, CF):
        sh = b.sheets[sheet]
        for (r, c), cell in sh.cells.items():
            if cell.period is not None and cell.period >= nh:
                assert cell.is_formula or cell.style == "input", f"{sheet}!{cell.address} is a hard-coded non-input"
    # the only hard-coded numbers in the Assumptions sheet are inputs
    for (r, c), cell in b.sheets["Assumptions"].cells.items():
        if not cell.is_formula and isinstance(cell.content, (int, float)):
            assert cell.style == "input"


def test_overrides_and_years(dataset):
    A = derive_assumptions(dataset, 7, {"terminal_growth": 0.02, "rev_growth": [0.10], "erp": 0.06, "projection_years": 3})
    assert A.years == 3
    assert A.values["rev_growth"] == [0.10, 0.10, 0.10]
    assert "erp" in A.overridden
    res = build_model(dataset, A)
    assert res.book.meta["np"] == 3
    assert abs(res.book.num(IS, "revenue", res.book.meta["nh"]) - res.book.num(IS, "revenue", res.book.meta["nh"] - 1) * 1.10) < 1e-6
    with pytest.raises(ValueError):
        derive_assumptions(dataset, 5, {"tax_rate": 0.99})
    res10 = build_model(dataset, years=10)
    assert res10.book.meta["np"] == 10 and res10.summary["error_cells"] == {}


def test_sensitivity_centre_matches_dcf(dataset):
    res = build_model(dataset, years=5)
    b = res.book
    centre = b.val("Sensitivity", "sg_2_2")
    assert abs(centre - b.num(DCF, "implied_price")) < 1e-6
    # exit-multiple table centre equals the exit method price
    ev_exit = b.num(DCF, "sum_pv") + b.num(DCF, "pv_tv_exit")
    price_exit = (ev_exit + b.num(DCF, "less_debt") + b.num(DCF, "plus_cash")) / b.num(DCF, "d_shares")
    assert abs(b.val("Sensitivity", "sm_2_2") - price_exit) < 1e-6


def test_workbook_writes(dataset):
    res = build_model(dataset, years=5)
    data = write_workbook(res.book)
    assert data[:2] == b"PK" and len(data) > 20000
    from openpyxl import load_workbook
    import io
    wb = load_workbook(io.BytesIO(data))
    assert wb.sheetnames == SHEET_ORDER
    ws = wb["Income Statement"]
    r, c = res.book.lookup(IS, "revenue", res.book.meta["nh"])
    assert str(ws.cell(row=r, column=c).value).startswith("=")
    assert "ImpliedSharePrice" in wb.defined_names


@pytest.mark.skipif(not soffice_available(), reason="LibreOffice not installed")
def test_libreoffice_recalculation_matches_engine(dataset):
    res = build_model(dataset, years=5)
    data = write_workbook(res.book)
    v = verify_workbook(res.book, data)
    assert v["status"] == "verified", v["mismatches"][:5]
    assert v["cells_checked"] > 1000
