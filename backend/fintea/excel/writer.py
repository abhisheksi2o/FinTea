"""Render a Book to a formatted .xlsx with openpyxl (formulas, not values).

Formatting follows investment-banking convention: blue inputs on a pale
yellow fill, black formulas, green cross-sheet links, bold totals with a top
border, Arial 10, frozen headers, and number formats that show negatives in
parentheses and zeros as dashes.
"""
from __future__ import annotations

import io
from typing import Optional

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName

from ..sheet import Book, Err
from ..sheet.book import FORMATS

FONT = "Arial"
BLUE, BLACK, GREEN, WHITE, GREY, RED, NAVY = "0000FF", "000000", "008000", "FFFFFF", "595959", "C00000", "1F3864"
INPUT_FILL = PatternFill("solid", fgColor="FFFFCC")
HEADER_FILL = PatternFill("solid", fgColor=NAVY)
SECTION_FILL = PatternFill("solid", fgColor="D9E1F2")
CHECK_FILL = PatternFill("solid", fgColor="E2EFDA")
THIN = Side(style="thin", color="7F7F7F")
MED = Side(style="medium", color=BLACK)


def _font(color=BLACK, bold=False, italic=False, size=10) -> Font:
    return Font(name=FONT, size=size, color=color, bold=bold, italic=italic)


def _apply_style(c, style: str, bold: bool, italic: bool, indent: int, wrap: bool, is_formula: bool, is_text: bool):
    if style == "input":
        c.font = _font(BLUE, bold)
        if not is_text:
            c.fill = INPUT_FILL
        else:
            c.fill = INPUT_FILL
    elif style == "link":
        c.font = _font(GREEN, bold)
    elif style == "total":
        c.font = _font(BLACK, True)
        if not is_text:
            c.border = Border(top=THIN)
    elif style == "header":
        c.font = _font(WHITE, True)
        c.fill = HEADER_FILL
        c.alignment = Alignment(horizontal="center" if not is_text or c.column > 1 else "left", vertical="center")
    elif style == "section":
        c.font = _font(NAVY, True, size=11)
        c.fill = SECTION_FILL
    elif style == "title":
        c.font = _font(NAVY, True, size=16)
    elif style == "subtitle":
        c.font = _font(GREY, bold, True, size=10)
    elif style == "note":
        c.font = _font(GREY, False, True, size=9)
    elif style == "memo":
        c.font = _font(GREY, False, True)
    elif style == "check":
        c.font = _font(RED, True)
        c.fill = CHECK_FILL
    elif style == "flag":
        c.font = _font(RED, True)
    else:  # formula / label / text
        c.font = _font(BLACK, bold, italic)
    if indent and c.column == 1:
        c.alignment = Alignment(indent=indent, wrap_text=wrap)
    elif wrap:
        c.alignment = Alignment(wrap_text=True, vertical="top")


def write_workbook(book: Book, path: Optional[str] = None) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    for name in book.order:
        sh = book.sheets[name]
        ws = wb.create_sheet(title=name)
        if sh.tab_color:
            ws.sheet_properties.tabColor = sh.tab_color
        ws.sheet_view.showGridLines = False
        for (r, col), cell in sh.cells.items():
            c = ws.cell(row=r, column=col)
            if cell.is_formula:
                c.value = book.formula(cell)
            else:
                v = cell.content
                if isinstance(v, Err):
                    v = v.code
                c.value = v
            is_text = isinstance(cell.content, str)
            c.number_format = FORMATS.get(cell.fmt, "General")
            _apply_style(c, cell.style, cell.bold, cell.italic, cell.indent, cell.wrap, cell.is_formula, is_text)
            if cell.hyperlink:
                c.hyperlink = cell.hyperlink
                c.font = Font(name=FONT, size=10, color="0563C1", underline="single")
        for col, w in sh.col_widths.items():
            ws.column_dimensions[get_column_letter(col)].width = w
        if not sh.col_widths:
            ws.column_dimensions["A"].width = 44
        for col in range(1, sh.max_col + 1):
            if col not in sh.col_widths and col > 1:
                ws.column_dimensions[get_column_letter(col)].width = 13
        for rng in sh.merges:
            ws.merge_cells(rng)
        for r, h in sh.row_heights.items():
            ws.row_dimensions[r].height = h
        if sh.freeze:
            ws.freeze_panes = sh.freeze
        ws.sheet_properties.outlinePr.summaryBelow = False
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth = 1
        ws.sheet_properties.pageSetUpPr.fitToPage = True
    # defined names for the headline outputs so the workbook is easy to audit
    for nm, (sheet, key) in {"WACC": ("WACC", "wacc"), "ImpliedSharePrice": ("DCF", "implied_price"),
                             "EnterpriseValue": ("DCF", "enterprise_value"), "SelectedBeta": ("Beta", "selected_beta"),
                             "TerminalGrowth": ("Assumptions", "terminal_growth"), "TaxRate": ("Assumptions", "tax_rate")}.items():
        if book.has(sheet, key):
            r, c = book.lookup(sheet, key, None)
            ref = f"'{sheet}'!${get_column_letter(c)}${r}"
            wb.defined_names[nm] = DefinedName(nm, attr_text=ref)
    wb.calculation.fullCalcOnLoad = True
    wb.properties.creator = "FinTea"
    wb.properties.title = f"{book.meta.get('company', '')} financial model"
    buf = io.BytesIO()
    wb.save(buf)
    data = buf.getvalue()
    if path:
        with open(path, "wb") as f:
            f.write(data)
    return data
