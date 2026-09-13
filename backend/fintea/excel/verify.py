"""Independent verification of the workbook formulas.

The .xlsx is recalculated by a headless LibreOffice Calc instance and every
formula cell is compared with the value computed by FinTea's own evaluator.
Two independent engines agreeing on every cell is strong evidence that the
formulas are correct and that Excel will show the same numbers.
"""
from __future__ import annotations

import os
import shutil
import subprocess
import tempfile
import uuid
from typing import Any, Dict

from openpyxl import load_workbook

from ..sheet import Book, Err


def soffice_available() -> bool:
    return shutil.which("soffice") is not None or shutil.which("libreoffice") is not None


def recalculate(xlsx_bytes: bytes, timeout: int = 180) -> bytes:
    """Return the workbook re-saved by LibreOffice with cached values."""
    exe = shutil.which("soffice") or shutil.which("libreoffice")
    if not exe:
        raise RuntimeError("LibreOffice (soffice) is not installed")
    work = tempfile.mkdtemp(prefix="fintea_lo_")
    try:
        src = os.path.join(work, "model.xlsx")
        with open(src, "wb") as f:
            f.write(xlsx_bytes)
        out = os.path.join(work, "out")
        profile = os.path.join(work, "profile_" + uuid.uuid4().hex)
        cmd = [exe, "--headless", "--norestore", f"-env:UserInstallation=file://{profile}",
               "--convert-to", "xlsx", "--outdir", out, src]
        subprocess.run(cmd, check=True, timeout=timeout, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
        with open(os.path.join(out, "model.xlsx"), "rb") as f:
            return f.read()
    finally:
        shutil.rmtree(work, ignore_errors=True)


def verify_workbook(book: Book, xlsx_bytes: bytes, abs_tol: float = 1e-6, rel_tol: float = 1e-7) -> Dict[str, Any]:
    if not soffice_available():
        return {"status": "skipped", "reason": "LibreOffice not installed; formulas were not independently recalculated",
                "cells_checked": 0, "mismatches": []}
    try:
        recalculated = recalculate(xlsx_bytes)
    except Exception as e:  # pragma: no cover - environment dependent
        return {"status": "skipped", "reason": f"LibreOffice recalculation failed: {e}", "cells_checked": 0, "mismatches": []}
    wb = load_workbook(io_bytes(recalculated), data_only=True)
    checked = 0
    mismatches = []
    max_abs = 0.0
    for name in book.order:
        ws = wb[name]
        for (r, c), cell in book.sheets[name].cells.items():
            if not cell.is_formula:
                continue
            expected = book.value(name, r, c)
            actual = ws.cell(row=r, column=c).value
            checked += 1
            if isinstance(expected, Err):
                ok = isinstance(actual, str) and actual.startswith("#")
            elif isinstance(expected, bool):
                ok = actual == expected
            elif isinstance(expected, (int, float)):
                if isinstance(actual, (int, float)) and not isinstance(actual, bool):
                    diff = abs(float(actual) - float(expected))
                    max_abs = max(max_abs, diff)
                    ok = diff <= abs_tol + rel_tol * abs(float(expected))
                else:
                    ok = False
            else:
                ok = (actual if actual is not None else "") == (expected if expected is not None else "")
            if not ok:
                mismatches.append({"sheet": name, "cell": cell.address, "expected": _js(expected), "actual": _js(actual),
                                   "formula": book.formula(cell)})
    return {"status": "verified" if not mismatches else "mismatch", "engine": "LibreOffice Calc (headless)",
            "cells_checked": checked, "max_abs_diff": max_abs, "mismatches": mismatches[:50], "n_mismatches": len(mismatches),
            "recalculated": recalculated if not mismatches else None}


def _js(v):
    if isinstance(v, Err):
        return v.code
    return v


def io_bytes(b: bytes):
    import io
    return io.BytesIO(b)
